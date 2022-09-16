"""
Use the instrument risk registers in the curren folder to update the databse for this month.

By default, possible information from a run within the last two weeks is being deleted to allow multiple
runs if some risk registers were delivered too late etc.

Written by Artur Glavic (artur.glavic@psi.ch).
"""

import logging
import os
import sys
from datetime import datetime, timedelta
from glob import glob

import pyodbc
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table


if getattr(sys, 'frozen', False):
    # If the application is run as a bundle, the PyInstaller bootloader
    # extends the sys module by a flag frozen=True and sets the app
    # path into variable _MEIPASS'.
    CUR_PATH = os.path.dirname(os.path.abspath(sys.executable))
else:
    CUR_PATH = os.path.dirname(os.path.abspath(__file__))

DATE_OVERWRITE = None  # datetime(year=2022, month=7, day=27)
RM_TABLE = str.maketrans(dict.fromkeys('aeiouAEIOU-_ '))


def handle_exception(exc_type, exc_value, exc_traceback):
    if issubclass(exc_type, KeyboardInterrupt):
        sys.__excepthook__(exc_type, exc_value, exc_traceback)
        return
    logging.error("Uncaught exception", exc_info=(exc_type, exc_value, exc_traceback))
    if getattr(sys, 'frozen', False):
        # keep console window open when run from .exe
        input("")


def setup_logging():
    logging.basicConfig(filename=os.path.join(CUR_PATH, 'update.log'), encoding='utf-8',
                        level=logging.DEBUG, format='%(levelname)-8s %(message)s', filemode='w')
    root = logging.getLogger()
    console = logging.StreamHandler()
    console.setLevel(logging.INFO)
    # set a format which is simpler for console use
    formatter = logging.Formatter('%(message)s')
    console.setFormatter(formatter)
    # add the handler to the root logger
    root.addHandler(console)
    sys.excepthook = handle_exception
    logging.debug('Started logging for update_databse script, (artur.glavic@psi.ch)')


def read_excel(fname):
    logging.info(f'Read risk data from {fname}')
    wb = load_workbook(filename=fname, data_only=True)
    tbl: Table = wb['Risks'].tables['Risk_Reg']
    data = [[ci.value for ci in ri] for ri in wb['Risks'][tbl.ref]]
    logging.debug(f'  found {len(data)-tbl.headerRowCount} rows')
    return tbl.column_names, data[tbl.headerRowCount:]


def tint(val):
    try:
        return int(val)
    except ValueError:
        return -1


def tdate(val):
    if isinstance(val, datetime):
        return val.date()
    else:
        return None


def check_row_entries(di):
    problem_columns = []
    try:
        di[0].upper()
    except (AttributeError, ValueError, TypeError):
        problem_columns.append(0)
    for i in [1, 17, 18, 19, 20, 21, 22, 24]:
        try:
            tint(di[i])
        except (AttributeError, ValueError, TypeError):
            problem_columns.append(i)
    for i in [14, 16]:
        try:
            tdate(di[i])
        except (AttributeError, ValueError, TypeError):
            problem_columns.append(i)
    logging.debug(f'    issues in columns {problem_columns}')


def main():
    setup_logging()
    logging.info(f'Database location: {CUR_PATH}\\full_risk_database.accdb')
    conn_str = (
        r'DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};'
        f'DBQ={CUR_PATH}\\full_risk_database.accdb;'
    )
    cnxn = pyodbc.connect(conn_str)
    crsr = cnxn.cursor()

    columns = [c.column_name for c in crsr.columns(table="Full Risk History")]
    logging.debug(f'  Columns in databse: {columns}')

    # delete entries younger than 2 weeks
    now_date = DATE_OVERWRITE or datetime.now().date()
    ref_data = now_date-timedelta(days=14)
    logging.debug('Deleting entreis from last 2 weeks and latest risks table')
    crsr.execute('DELETE FROM "Full Risk History" WHERE DateValue("Date Added")>=?', ref_data)
    crsr.execute('DELETE FROM "Latest Risks"')

    for fname in glob(os.path.join(CUR_PATH, 'latest', '*Risks.xlsx')):
        dcols, data = read_excel(fname)

        insert_data = []

        for i, di in enumerate(data):
            try:
                inst = di[0].upper().translate(RM_TABLE)[:3]
                if len(inst)<3:
                    inst = di[0].upper()[:3]
                row = (now_date, di[0].upper(), tint(di[1]), f'{inst}-{tint(di[1]):02}', di[2], di[3],
                       di[4], di[5], di[6], di[7], di[8],
                       tdate(di[14]), di[15], tdate(di[16]),
                       tint(di[17]), tint(di[18]), tint(di[19]), tint(di[20]), tint(di[21]), tint(di[22]), tint(di[24])
                       )
            except (AttributeError, ValueError, TypeError):
                logging.warning(f'  error when parsing row {i+5}, check table')
                logging.debug(f'    data: {di}')
                check_row_entries(di)
                logging.debug('    error message:', exc_info=True)
                continue

            insert_data.append(row)

            res = crsr.execute('SELECT "Risk Rating" FROM "Full Risk History" WHERE Project=? AND "Risk ID"=? '
                               'ORDER BY "Date Added" DESC',
                               (row[1], row[2]))
            risk_history = [fi[0] for fi in res.fetchall()]
            try:
                prev_rating = risk_history[0]
            except IndexError:
                prev_rating = -1
            crsr.execute('INSERT INTO "Latest Risks"'
                         '(Project, "Risk ID", "Global ID", "Risk Title", "Risk and Impact Description",'
                         'Owner, Partner, Status, "Risk Treatment", "Past Treatment Actions and Notes", '
                         '"Last Reviewed", "Planned Treatment Actions", "Action Due", '
                         'when, cost, schedule, quality, "max impact", likelihood, "Risk Rating", '
                         '"Last Rating", "Full History") '
                         'VALUES (?, ?, ?, ?, ?, '
                         '?, ?, ?, ?, ?, '
                         '?, ?, ?, '
                         '?, ?, ?, ?, ?, ?, ?,'
                         '?, ?)',
                         row[1:]+(prev_rating, repr(risk_history)))

        logging.debug(f'  inserting {len(insert_data)} lines into databse')
        crsr.executemany('INSERT INTO "Full Risk History"'
                         '("Date Added", Project, "Risk ID", "Global ID", "Risk Title", "Risk and Impact Description",'
                         'Owner, Partner, Status, "Risk Treatment", "Past Treatment Actions and Notes", '
                         '"Last Reviewed", "Planned Treatment Actions", "Action Due", '
                         'when, cost, schedule, quality, "max impact", likelihood, "Risk Rating") '
                         'VALUES (?, ?, ?, ?, ?, ?, '
                         '?, ?, ?, ?, ?, '
                         '?, ?, ?, '
                         '?, ?, ?, ?, ?, ?, ?)',
                         insert_data)
    crsr.commit()

    if getattr(sys, 'frozen', False):
        # keep console window open when run from .exe
        input("Finished, press enter to close program.")


if __name__=='__main__':
    main()
